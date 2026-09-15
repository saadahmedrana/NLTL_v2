#!/usr/bin/env python3
from __future__ import annotations

import csv
import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from rdflib import Graph, RDF, URIRef
from rdflib.namespace import SH

import enrich_automatic_evidence as enrich


SHAPE_TTL = """@prefix sh: <http://www.w3.org/ns/shacl#> .
@prefix ex: <https://example.test/> .
ex:ShipShape a sh:NodeShape ;
  sh:targetClass ex:Ship ;
  sh:property [ a sh:PropertyShape ; sh:path ex:name ; sh:minCount 1 ] .
"""
PASS_TTL = """@prefix ex: <https://example.test/> .
ex:s a ex:Ship ; ex:name "Ship" .
"""
FAIL_TTL = """@prefix ex: <https://example.test/> .
ex:s a ex:Ship .
"""


class AutomaticEvidenceUnitTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.allowed = enrich.load_supported_shacl_vocabulary()

    def test_hash_validation(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "x.ttl"
            path.write_text("abc", encoding="utf-8")
            self.assertEqual(enrich.sha256_file(path), hashlib.sha256(b"abc").hexdigest())

    def test_meta_shacl_derives_only_from_successful_exact_execution(self):
        sha = "a" * 64
        row = {
            "generated_shape_sha256": sha,
            "shape_parse_status": "PARSED",
            "execution_status": "EXECUTED",
            "pyshacl_options": {"meta_shacl": True},
        }
        self.assertEqual(
            enrich.derive_meta_shacl_from_rows([row], sha),
            ("VALID", "DERIVED_FROM_SUCCESSFUL_META_SHACL_EXECUTION"),
        )
        for field, value in (
            ("generated_shape_sha256", "b" * 64),
            ("shape_parse_status", "NOT_PARSED"),
            ("execution_status", "INFRASTRUCTURE_FAILURE"),
            ("pyshacl_options", {"meta_shacl": False}),
        ):
            changed = dict(row)
            changed[field] = value
            self.assertIsNone(enrich.derive_meta_shacl_from_rows([changed], sha))

    def test_vocabulary_check_ignores_project_terms_and_rejects_unknown_shacl(self):
        valid = Graph().parse(data=SHAPE_TTL, format="turtle")
        self.assertEqual(enrich.check_shacl_vocabulary(valid, self.allowed), ("VALID", []))
        invalid = Graph().parse(
            data=SHAPE_TTL + "<https://example.test/x> sh:minCounnt 2 .\n",
            format="turtle",
        )
        status, unknown = enrich.check_shacl_vocabulary(invalid, self.allowed)
        self.assertEqual(status, "INVALID")
        self.assertEqual(unknown, [str(SH) + "minCounnt"])

    def test_standard_targets_activate(self):
        data = Graph().parse(data="""@prefix ex: <https://example.test/> .
            ex:s a ex:Ship ; ex:p ex:o . ex:Sub a <http://www.w3.org/2000/01/rdf-schema#Class> ;
            <http://www.w3.org/2000/01/rdf-schema#subClassOf> ex:Ship . ex:t a ex:Sub .""", format="turtle")
        for predicate, obj in (
            ("targetNode", "ex:s"), ("targetClass", "ex:Ship"),
            ("targetSubjectsOf", "ex:p"), ("targetObjectsOf", "ex:p"),
        ):
            shapes = Graph().parse(data=f"""@prefix sh: <{SH}> . @prefix ex: <https://example.test/> .
                ex:S a sh:NodeShape ; sh:{predicate} {obj} .""", format="turtle")
            result, nodes, _ = enrich.target_activation(shapes, data)
            self.assertEqual(result, "ACTIVATED")
            self.assertTrue(nodes)

    def test_not_activated_is_not_inferred_from_report_focus_nodes(self):
        shapes = Graph().parse(data=SHAPE_TTL, format="turtle")
        data = Graph().parse(data="@prefix ex: <https://example.test/> . ex:x ex:p ex:y .", format="turtle")
        self.assertEqual(enrich.target_activation(shapes, data)[0], "NOT_ACTIVATED")

    def test_custom_target_is_not_evaluated(self):
        shapes = Graph().parse(data=f"""@prefix sh: <{SH}> . @prefix ex: <https://example.test/> .
            ex:S a sh:NodeShape ; sh:target [ a sh:SPARQLTarget ; sh:select "SELECT ?this WHERE {{ ?this ?p ?o }}" ] .""", format="turtle")
        result, _, warnings = enrich.target_activation(shapes, Graph())
        self.assertEqual(result, "NOT_EVALUATED")
        self.assertIn("custom", warnings[0])

    def test_node_and_property_shape_counts(self):
        profile = enrich.structural_profile(Graph().parse(data=SHAPE_TTL, format="turtle"))
        self.assertEqual(profile["node_shape_count"], 1)
        self.assertEqual(profile["property_shape_count"], 1)

    def test_core_component_counting_excludes_metadata_targets_and_sparql(self):
        graph = Graph().parse(data=f"""@prefix sh: <{SH}> . @prefix ex: <https://example.test/> .
            ex:S a sh:NodeShape ; sh:targetClass ex:C ; sh:message "m" ; sh:class ex:C ;
              sh:sparql [ sh:select "SELECT ?this WHERE {{ ?this ?p ?o }}" ] .""", format="turtle")
        profile = enrich.structural_profile(graph)
        self.assertEqual(profile["shacl_core_component_occurrence_count"], 1)
        self.assertEqual(profile["shacl_core_components_used"], ["sh:class"])
        self.assertEqual(profile["sparql_constraint_count"], 1)

    def test_path_depth(self):
        graph = Graph().parse(data=f"""@prefix sh: <{SH}> . @prefix ex: <https://example.test/> .
            ex:P sh:path ( ex:p [ sh:inversePath ex:q ] ) .""", format="turtle")
        path = next(graph.objects(URIRef("https://example.test/P"), SH.path))
        self.assertEqual(enrich.shacl_path_depth(graph, path), 3)

    def test_logical_nesting(self):
        graph = Graph().parse(data=f"""@prefix sh: <{SH}> . @prefix ex: <https://example.test/> .
            ex:S sh:or ( [ sh:not [ sh:class ex:C ] ] ) .""", format="turtle")
        self.assertEqual(enrich.logical_nesting_depth(graph, URIRef("https://example.test/S")), 2)

    def test_ledger_reconciliation_prevents_artifact_substitution(self):
        record = {
            "architecture": "NO_SEMANTIC", "requirement_id": "R1",
            "selected_case_ids": ["R1-P", "R1-F"], "shape_sha256": "a" * 64,
        }
        rows = [{
            "configuration": "NO_SEMANTIC", "requirement_id": "R1", "case_id": case,
            "generation_run": "RUN_01", "generated_shape_sha256": "b" * 64,
        } for case in record["selected_case_ids"]]
        matched, discrepancies = enrich.reconcile_ledger_rows(record, rows, True)
        self.assertEqual(matched, [])
        self.assertEqual(len(discrepancies), 2)

    def test_missing_artifact_handling(self):
        with tempfile.TemporaryDirectory() as temporary:
            repo = Path(temporary)
            ledger = repo / "ledger.jsonl"
            rows = [self._ledger_row("R1-P", None), self._ledger_row("R1-F", None)]
            ledger.write_text("".join(json.dumps(row) + "\n" for row in rows), encoding="utf-8")
            record = self._selection_record("V2_FALLBACK25", "R1", None, None, "ledger.jsonl", enrich.sha256_file(ledger))
            evidence, _ = enrich.build_evidence_record(repo, record, enrich.LedgerCache(repo), self.allowed, False)
            self.assertEqual(evidence["availability"], "MISSING_ARTIFACT")
            self.assertEqual(evidence["shacl_vocabulary_validity"], "MISSING_ARTIFACT")

    def test_required_workbook_headers_are_discovered_by_text(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "book.xlsx"
            self._make_workbook(path, ["R1"])
            book = enrich.WorkbookXML(path)
            row, headers = book.discover_headers("V2", enrich.REQUIRED_HEADERS)
            self.assertEqual(row, 1)
            self.assertEqual(headers["Requirement ID"], 1)
            with self.assertRaises(enrich.EnrichmentError):
                book.discover_headers("V2", (*enrich.REQUIRED_HEADERS, "Absent MethodV2 field"))

    def test_synthetic_end_to_end_preserves_manual_cells_formulas_and_features(self):
        with tempfile.TemporaryDirectory() as temporary:
            repo = Path(temporary)
            audit = repo / "audit"
            audit.mkdir()
            input_book = audit / enrich.INPUT_WORKBOOK_NAME
            self._make_workbook(input_book, ["R1"])
            rdf_root = repo / "MVP/SHACL_GENERATION_PIPELINE/evaluation/BEHAVIORAL_RDF_R13/rdf/SYN"
            rdf_root.mkdir(parents=True)
            (rdf_root / "R1-P.ttl").write_text(PASS_TTL, encoding="utf-8")
            (rdf_root / "R1-F.ttl").write_text(FAIL_TTL, encoding="utf-8")
            artifact_root = repo / "artifacts"
            artifact_root.mkdir()
            records = []
            ledger_rows = []
            for architecture in enrich.ARCHITECTURES:
                shape = artifact_root / f"{architecture}.ttl"
                shape.write_text(SHAPE_TTL, encoding="utf-8")
                shape_sha = enrich.sha256_file(shape)
                for case_id, expected, conforms in (("R1-P", "PASS", True), ("R1-F", "FAIL", False)):
                    rdf_path = rdf_root / f"{case_id}.ttl"
                    ledger_rows.append(self._ledger_row(case_id, shape_sha, architecture, expected, conforms, rdf_path))
                records.append(self._selection_record(
                    architecture, "R1", str(shape.relative_to(repo)), shape_sha,
                    "ledger.jsonl", "PENDING",
                ))
            ledger = repo / "ledger.jsonl"
            ledger.write_text("".join(json.dumps(row) + "\n" for row in ledger_rows), encoding="utf-8")
            ledger_sha = enrich.sha256_file(ledger)
            for record in records:
                record["source_ledger_sha256"] = ledger_sha
            selection = {
                "requirements": ["R1"], "records": records,
                "baseline_metrics": {
                    architecture: {"cases": 2, "correct": 2, "executed": 2, "macro": 1.0, "exact": 1, "requirements": 1, "fa": 0, "fr": 0}
                    for architecture in enrich.ARCHITECTURES
                },
            }
            (audit / enrich.SELECTION_NAME).write_text(json.dumps(selection), encoding="utf-8")
            result = enrich.run_enrichment(
                repo, audit, expected_requirements=1,
                expected_availability={architecture: 1 for architecture in enrich.ARCHITECTURES},
                recheck_unresolved=False,
            )
            self.assertEqual(result["records"], 3)
            verification = enrich.verify_outputs(repo, audit, expected_requirements=1)
            self.assertEqual(verification["status"], "PASS")
            workbook = load_workbook(audit / enrich.OUTPUT_WORKBOOK_NAME, data_only=False)
            for sheet_name in enrich.SHEET_BY_ARCHITECTURE.values():
                sheet = workbook[sheet_name]
                headers = {cell.value: cell.column for cell in sheet[1]}
                self.assertEqual(sheet.cell(2, headers["Manual field"]).value, "HUMAN_KEEP")
                self.assertEqual(sheet.cell(2, headers["Manual formula"]).value, "=1+1")
                self.assertEqual(sheet.cell(2, headers["SHACL specification validity"]).value, "PASS")
                self.assertEqual(sheet.cell(2, headers["SHACL vocabulary validity"]).value, "VALID")
                self.assertIn("ACTIVATED", sheet.cell(2, headers["Selected-case target activation"]).value)
                self.assertTrue(sheet.cell(2, headers["Manual field"]).comment)
                self.assertTrue(sheet.cell(2, headers["Manual field"]).hyperlink)
                self.assertEqual(len(sheet.data_validations.dataValidation), 1)
            workbook.close()

    @staticmethod
    def _ledger_row(
        case_id: str,
        shape_sha: str | None,
        architecture: str = "V2_FALLBACK25",
        expected: str | None = None,
        conforms: bool | None = None,
        rdf_path: Path | None = None,
    ) -> dict:
        configuration = {
            "V2_FALLBACK25": "FULL_REPAIR_V2",
            "NO_SEMANTIC": "NO_SEMANTIC",
            "SINGLESHOT": "SINGLESHOT",
        }[architecture]
        expected = expected or ("PASS" if case_id.endswith("P") else "FAIL")
        return {
            "configuration": configuration, "requirement_id": "R1", "case_id": case_id,
            "generation_run": "RUN_01", "generated_shape_sha256": shape_sha,
            "rdf_path": f"rdf/SYN/{case_id}.ttl",
            "rdf_sha256": enrich.sha256_file(rdf_path) if rdf_path else None,
            "expected_outcome": expected, "actual_conforms": conforms,
            "execution_status": "EXECUTED" if shape_sha else "INFRASTRUCTURE_FAILURE",
            "behavioral_match": conforms == (expected == "PASS") if conforms is not None else None,
            "outcome_class": "CORRECT_BEHAVIOR" if conforms == (expected == "PASS") else "GENERATION_ERROR",
            "shape_parse_status": "PARSED" if shape_sha else "NOT_PARSED",
            "generation_pipeline_status": "GENERATION_ACCEPTED" if shape_sha else "GENERATION_REJECTED",
            "pyshacl_options": {"meta_shacl": True}, "pyshacl_version": "synthetic", "rdflib_version": "synthetic",
        }

    @staticmethod
    def _selection_record(
        architecture: str,
        requirement_id: str,
        shape_path: str | None,
        shape_sha: str | None,
        ledger_path: str,
        ledger_sha: str,
    ) -> dict:
        return {
            "architecture": architecture, "requirement_id": requirement_id,
            "generation_run": "RUN_01", "selected_generation_run": "RUN_01",
            "shape_path": shape_path, "shape_sha256": shape_sha,
            "output_origin": "OFFICIAL_ELIGIBLE_OUTPUT" if shape_path else "NO_ELIGIBLE_OUTPUT",
            "parse_status": "PARSED" if shape_path else "NOT_PARSED",
            "deterministic_status": "PASS" if shape_path else "FAIL",
            "original_generation_status": "GENERATION_ACCEPTED" if shape_path else "GENERATION_REJECTED",
            "selected_case_ids": ["R1-P", "R1-F"],
            "source_ledger_path": ledger_path, "source_ledger_sha256": ledger_sha,
            "deterministic_validation_path": None, "deterministic_validation_sha256": None,
            "artifact_metadata_path": None, "candidate_attempt": 1,
            "deterministic_findings": {},
        }

    @staticmethod
    def _make_workbook(path: Path, requirements: list[str]) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        headers = [
            "Requirement ID", "Selected generation run", "Selected SHACL path", "Selected SHACL SHA-256",
            "Recorded Turtle parse status", "Recorded deterministic status", "Selected case 1",
            "Original expected 1", "Recorded actual 1", "Recorded outcome 1", "Selected case 2",
            "Original expected 2", "Recorded actual 2", "Recorded outcome 2", *enrich.AUTOMATIC_HEADERS,
            "Manual field", "Manual formula",
        ]
        for architecture, sheet_name in enrich.SHEET_BY_ARCHITECTURE.items():
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(headers)
            for requirement_id in requirements:
                row = [
                    requirement_id, "RUN_01", f"artifacts/{architecture}.ttl", "PENDING", "PARSED", "PASS",
                    f"{requirement_id}-P", "PASS", "PASS", "CORRECT_BEHAVIOR",
                    f"{requirement_id}-F", "FAIL", "FAIL", "CORRECT_BEHAVIOR",
                    *([None] * len(enrich.AUTOMATIC_HEADERS)), "HUMAN_KEEP", "=1+1",
                ]
                sheet.append(row)
                manual = sheet.cell(sheet.max_row, headers.index("Manual field") + 1)
                manual.comment = Comment("Keep this decision", "Reviewer")
                manual.hyperlink = "https://example.test/evidence"
            validation = DataValidation(type="list", formula1='"HUMAN_KEEP,OTHER"')
            sheet.add_data_validation(validation)
            validation.add(sheet.cell(2, headers.index("Manual field") + 1))
        results = workbook.create_sheet("Results")
        results["A1"] = "Baseline"
        results["B1"] = 2
        workbook.save(path)


if __name__ == "__main__":
    unittest.main(verbosity=2)
