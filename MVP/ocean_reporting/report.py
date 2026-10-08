#!/usr/bin/env python3
"""Offline reporting only: reads saved evidence; never executes SHACL or calls APIs."""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import itertools
import json
import math
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd

csv.field_size_limit(100_000_000)
HERE = Path(__file__).resolve().parent
PIPE = "SHACL_GENERATION_PIPELINE"
SCALING = f"{PIPE}/experiments/MODEL_SCALING_FULL_REPAIR_V2_50"
EVALUATIONS = f"{PIPE}/evaluation/experiment_results"
WORKFLOW_RUNS = {
    "official": "BEHAVIORAL-R13-FULL-REPAIR-V2-RUN01-20260915T144116505813Z",
    "historical": "BEHAVIORAL-R13-20260911T080233539977Z",
    "rejected": "BEHAVIORAL-R13-FULL-REPAIR-V2-REJECTED-A4-RUN01-20260915T155743340362Z",
}
MODELS = ["luna", "sol", "gemini", "gpt_oss"]
STAGES = ["FIRST_GENERATION", "SELF_REPAIR_FINAL"]
WORKFLOWS = ["FULL_REPAIR_V2", "NO_SEMANTIC", "SINGLESHOT", "RETROSPECTIVE_V2_PLUS_REJECTED25"]
COLORS = dict(zip(MODELS + WORKFLOWS, ["#0072B2", "#009E73", "#D55E00", "#CC79A7", "#0072B2", "#009E73", "#D55E00", "#777777"]))
VERDICTS = ["CORRECT", "MINOR_DEFECT", "MAJOR_DEFECT", "NOT_ASSESSABLE", "UNCLEAR"]
VERDICT_COLORS = ["#009E73", "#E69F00", "#D55E00", "#999999", "#CC79A7"]


def digest(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def clean(value):
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, dict):
        return {str(k): clean(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean(v) for v in value]
    if isinstance(value, np.generic):
        return clean(value.item())
    if value is None or (isinstance(value, float) and not math.isfinite(value)):
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def write_json(path, value):
    Path(path).write_text(json.dumps(clean(value), indent=2, ensure_ascii=False, allow_nan=False), encoding="utf-8")


def export(out, name, records):
    frame = records if isinstance(records, pd.DataFrame) else pd.DataFrame(records)
    frame.to_csv(out / f"{name}.csv", index=False)
    write_json(out / f"{name}.json", frame.to_dict("records"))


def require(condition, message):
    if not condition:
        raise ValueError(message)


class Evidence:
    def __init__(self, root, output):
        self.root, self.output = root, output
        self.sources, self.references, self.missing, self.checks = [], [], [], []
        self.used = set()

    def use(self, path, role, identity, required=True):
        path = Path(path).resolve()
        if not path.is_file():
            item = {"path": str(path), "role": role, "identity": identity, "required": required}
            if item not in self.missing:
                self.missing.append(item)
            return False
        key = (str(path), role, identity)
        if key not in self.used:
            self.used.add(key)
            self.sources.append({"absolute_path": str(path), "role": role, "experiment_identity": identity, "sha256": digest(path)})
        return True

    def json(self, path, role, identity, required=True):
        return json.loads(Path(path).read_text(encoding="utf-8")) if self.use(path, role, identity, required) else None

    def csv(self, path, role, identity):
        if not self.use(path, role, identity):
            return []
        with Path(path).open(encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream))

    def jsonl(self, path, role, identity):
        if not self.use(path, role, identity):
            return []
        return [json.loads(line) for line in Path(path).read_text(encoding="utf-8").splitlines() if line.strip()]

    def reference(self, raw, base, role, identity):
        if not isinstance(raw, str) or not raw.strip():
            return None
        p = Path(raw)
        candidates = [p] if p.is_absolute() else [base / p, self.root / p, self.root.parent / p]
        if raw.startswith("MVP/"):
            candidates.insert(0, self.root / raw[4:])
        if raw.startswith("rdf/"):
            candidates.insert(0, self.root / PIPE / "evaluation/BEHAVIORAL_RDF_R13" / p)
        existing = next((x.resolve() for x in candidates if x.is_file()), None)
        self.references.append({"recorded_path": raw, "resolved_path": str(existing or candidates[0].resolve()),
                                "exists": existing is not None, "role": role, "experiment_identity": identity})
        return existing

    def collect_references(self, value, base, identity):
        if isinstance(value, dict):
            for key, child in value.items():
                if isinstance(child, str) and (key.endswith("path") or key.endswith("_file") or key == "source_ledger"):
                    self.reference(child, base, key, identity)
                elif isinstance(child, (dict, list)):
                    self.collect_references(child, base, identity)
        elif isinstance(value, list):
            for child in value:
                self.collect_references(child, base, identity)

    def check(self, name, condition, detail=""):
        self.checks.append({"check": name, "passed": bool(condition), "detail": detail})
        require(condition, f"{name}: {detail}")

    def save(self):
        export(self.output, "sources_used", self.sources)
        refs = pd.DataFrame(self.references)
        if not refs.empty:
            refs = refs.drop_duplicates()
        export(self.output, "referenced_files", refs)
        export(self.output, "verification", self.checks)
        lines = ["# Missing data", "", "Missing evaluation rows are distinct from missing generated artifacts.", ""]
        lines += [f"- {'Required' if r['required'] else 'Supporting'} {r['role']} ({r['identity']}): `{r['path']}`" for r in self.missing]
        if not self.missing:
            lines.append("No required input records are missing.")
        if not refs.empty:
            count = int((~refs.exists).sum())
            lines += ["", f"{count} referenced evidence paths could not be resolved. Exact paths and roles are in referenced_files.csv.",
                      "Unavailable references do not establish that a generated artifact was absent at evaluation time."]
        (self.output / "MISSING_DATA.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def workbook_data(evidence, workbook):
    require(evidence.use(workbook, "author manual reviews and case-review decisions", "manual audit"), f"Workbook missing: {workbook}")
    cached = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(workbook, read_only=True, data_only=False)
    sheets, cell_formulas = {}, []
    rawdir = evidence.output / "workbook_sheets"
    rawdir.mkdir()
    for sheet in cached:
        rows = list(sheet.values)
        sheets[sheet.title] = rows
        safe = re.sub(r"[^a-zA-Z0-9_-]", "_", sheet.title)
        with (rawdir / f"{safe}.csv").open("w", newline="", encoding="utf-8") as stream:
            writer = csv.writer(stream)
            writer.writerows([[clean(v) for v in row] for row in rows])
        write_json(rawdir / f"{safe}.json", {"sheet": sheet.title, "rows": rows})
        for row in formulas[sheet.title]:
            for cell in row:
                if cell.data_type == "f":
                    cell_formulas.append({"sheet": sheet.title, "cell": cell.coordinate, "formula": cell.value,
                                          "cached_value": rows[cell.row - 1][cell.column - 1]})
    export(evidence.output, "workbook_formulas_and_caches", cell_formulas)
    cached.close()
    formulas.close()
    return sheets


def table(rows, header_row, key):
    headers = rows[header_row - 1]
    return [dict(zip(headers, row)) | {"_excel_row": i} for i, row in enumerate(rows[header_row:], header_row + 1)
            if dict(zip(headers, row)).get(key)]


def load_benchmark(evidence, sheets):
    base = evidence.root / PIPE / "evaluation/BEHAVIORAL_RDF_R13/manifests"
    files = sorted(base.glob("*_batch_*_manifest.jsonl"))
    require(files, f"No frozen benchmark manifests: {base}")
    records = []
    for path in files:
        records += evidence.jsonl(path, "frozen benchmark case manifest", "R13 full benchmark")
    benchmark = pd.DataFrame(records)
    require({"case_id", "requirement_id", "expected"}.issubset(benchmark), "Benchmark schema missing IDs/expected labels")
    evidence.check("benchmark IDs unique", not benchmark.case_id.duplicated().any())
    evidence.check("original benchmark sizes", len(benchmark) == 2186 and benchmark.requirement_id.nunique() == 268)
    require(set(benchmark.expected) <= {"PASS", "FAIL"}, "Invalid benchmark expected verdict")
    metadata = {r["Requirement ID"]: r for r in table(sheets["V2"], 6, "Requirement ID")}
    benchmark["source_family"] = benchmark.requirement_id.map(lambda r: metadata[r]["Source family"])
    benchmark["verification_mode"] = benchmark.requirement_id.map(lambda r: metadata[r]["Verification mode"])
    benchmark["expected"] = benchmark.expected.astype(str)
    reviews = table(sheets["CaseReview"], 6, "Case ID")
    evidence.check("CaseReview case alignment", {r["Case ID"] for r in reviews} == set(benchmark.case_id))
    labels = dict(zip(benchmark.case_id, benchmark.expected))
    evidence.check("workbook original labels match frozen benchmark", all(r["Original expected verdict"] == labels[r["Case ID"]] for r in reviews))
    proposed = [r for r in reviews if str(r.get("Change type", "")).strip().upper() == "RDF_EDIT"]
    evidence.check("proposed RDF_EDIT count", len(proposed) == 26)
    export(evidence.output, "proposed_exclusions", proposed)
    return benchmark, {r["Case ID"] for r in proposed}


def selected_runs(evidence, overrides):
    runs = {}
    for role, default in WORKFLOW_RUNS.items():
        path = evidence.root / EVALUATIONS / default
        runs[role] = Path(overrides.get(role, path)).resolve()
    return runs


def normalise(evidence, rows, benchmark, condition, stage, base, generation):
    """Left join onto frozen cases. Missing ledger rows are explicit, never artifact claims."""
    bycase = {}
    for row in rows:
        cid = row.get("case_id")
        require(cid not in bycase, f"Duplicate case {cid} in {condition}/{stage}")
        bycase[cid] = row
    expected_ids = set(benchmark.case_id)
    require(set(bycase) <= expected_ids, f"Unexpected case IDs in {condition}/{stage}: {set(bycase) - expected_ids}")
    missing = expected_ids - set(bycase)
    if missing:
        evidence.missing.append({"path": str(base / "raw_case_results.csv"), "role": f"{len(missing)} missing outcome rows: {','.join(sorted(missing))}",
                                 "identity": f"{condition}/{stage}", "required": True})
    records = []
    for b in benchmark.to_dict("records"):
        row = bycase.get(b["case_id"], {})
        if row:
            require(row.get("requirement_id") == b["requirement_id"], f"Requirement mismatch for {b['case_id']}")
            require(row.get("expected_outcome") == b["expected"], f"Expected label mismatch for {b['case_id']}")
        value = str(row.get("actual_conforms", "")).lower()
        executed = row.get("execution_status") == "EXECUTED"
        require(not executed or value in {"true", "false"}, f"Executed case lacks boolean verdict: {b['case_id']}")
        observed = ("PASS" if value == "true" else "FAIL") if executed else "NO_VERDICT"
        outcome = "NO_VERDICT" if observed == "NO_VERDICT" else ("CORRECT" if observed == b["expected"] else "FALSE_ACCEPT" if observed == "PASS" else "FALSE_REJECT")
        gen = generation.get(b["requirement_id"], {})
        gpath = gen.get("generated_shacl_path") or row.get("generated_shape_path") or ""
        resolved = evidence.reference(gpath, base, "generated artifact", f"{condition}/{stage}")
        availability = "RECORDED_ARTIFACT" if gpath else "RECORDED_NO_OUTPUT" if gen else "UNKNOWN"
        diagnostic = {k: row.get(k, "") or gen.get(k, "") for k in ["failure_stage", "exception_type", "exception_message", "failure_detail", "generation_status", "shape_parse_status", "semantic_validator_feedback", "semantic_validator_pipeline_decision"]}
        reason = "MISSING_EVALUATION_ROW" if not row else diagnostic["failure_stage"] or diagnostic["exception_type"] or diagnostic["generation_status"] or "NO_VERDICT_REASON_UNAVAILABLE"
        if observed != "NO_VERDICT":
            reason = "BEHAVIORAL_MISMATCH_CAUSE_UNDETERMINED" if outcome != "CORRECT" else "CORRECT"
        record = dict(row)
        record.update({"condition": condition, "stage": stage, "requirement_id": b["requirement_id"], "case_id": b["case_id"],
                       "source_family": b["source_family"], "verification_mode": b["verification_mode"], "expected": b["expected"],
                       "observed": observed, "scored_outcome": outcome, "correct": int(outcome == "CORRECT"),
                       "evaluation_row_present": bool(row), "artifact_availability": availability, "artifact_exists_now": resolved is not None,
                       "failure_reason_evidence_label": reason, "diagnostics": json.dumps(clean(diagnostic), ensure_ascii=False),
                       "provenance_ledger": str(base / "raw_case_results.csv"), "retrospective_candidate": False})
        evidence.collect_references(row, base, f"{condition}/{stage}")
        records.append(record)
    return records


def genmap(evidence, path, identity):
    rows = evidence.jsonl(path, "generation artifact and diagnostic manifest", identity)
    require(len({r['requirement_id'] for r in rows}) == len(rows), f"Duplicate generation requirement: {path}")
    for row in rows:
        evidence.collect_references(row, path.parent, identity)
    export(evidence.output, "generation_" + re.sub(r"\W", "_", identity), rows)
    return {r["requirement_id"]: r for r in rows}


def load_workflows(evidence, benchmark, runs):
    all_rows, supplemental = [], []
    manifests = evidence.root / PIPE / "evaluation/generated_rule_manifests/RUN_01"
    for role, selections in [("official", ["FULL_REPAIR_V2"]), ("historical", ["NO_SEMANTIC", "SINGLESHOT"]),
                             ("rejected", ["FULL_REPAIR_V2_REJECTED_ATTEMPT4_DIAGNOSTIC"])]:
        base = runs[role]
        manifest = evidence.json(base / "run_manifest.json", "selected evaluation identity", role)
        require(manifest is not None, f"Missing workflow run manifest: {base}")
        manifest_id = manifest.get("run_id") or manifest.get("diagnostic_run_id")
        require(manifest_id == base.name, f"Workflow run ID mismatch: {base}")
        evidence.collect_references(manifest, base, role)
        rows = evidence.csv(base / "raw_case_results.csv", "selected behavioural outcome ledger", role)
        for condition in selections:
            selected = [r for r in rows if r.get("configuration") == condition]
            require(selected, f"Condition {condition} absent from {base}")
            require(all(r.get('run_id') == manifest_id for r in selected), f"Mixed evaluation runs: {base}")
            filename = {"FULL_REPAIR_V2": "full_repair_v2_generated_rules.jsonl", "NO_SEMANTIC": "no_semantic_generated_rules.jsonl",
                        "SINGLESHOT": "singleshot_generated_rules.jsonl", "FULL_REPAIR_V2_REJECTED_ATTEMPT4_DIAGNOSTIC": "full_repair_v2_rejected_attempt4_diagnostic.jsonl"}[condition]
            generation = genmap(evidence, manifests / filename, condition)
            if role != 'rejected':
                evidence.check(f'{condition} generation manifest coverage', set(generation) == set(benchmark.requirement_id))
            subset = benchmark if role != "rejected" else benchmark[benchmark.requirement_id.isin(generation)]
            normal = normalise(evidence, selected, subset, condition, "RUN_01", base, generation)
            if role == "rejected":
                evidence.check("rejected supplement size", len(generation) == 25 and len(normal) == 199)
                supplemental = normal
            else:
                all_rows += normal
    official = [r for r in all_rows if r['condition'] == 'FULL_REPAIR_V2']
    supplement_map = {r['case_id']: r for r in supplemental}
    retrospective = []
    for row in official:
        replacement = supplement_map.get(row['case_id'])
        if replacement:
            evidence.check(f"retrospective does not replace accepted output {row['case_id']}", row['observed'] == 'NO_VERDICT')
        retro = dict(replacement or row)
        retro.update(condition=WORKFLOWS[-1], retrospective_candidate=replacement is not None)
        retrospective.append(retro)
    export(evidence.output, 'rejected_candidate_outcomes', supplemental)
    return pd.DataFrame(all_rows + retrospective)


def load_models(evidence, benchmark, overrides):
    root = Path(overrides.get("scaling_root", evidence.root / SCALING)).resolve()
    sample = evidence.json(root / "MANIFESTS/sample_50.json", "frozen common model sample", "SCALING50")
    require(sample is not None, "Missing model sample manifest")
    require(len(sample['records']) == 50, "Model sample does not contain 50 requirements")
    selected = {r['requirement_id'] for r in sample['records']}
    subset = benchmark[benchmark.requirement_id.isin(selected)]
    evidence.check("model sample denominator agrees with frozen counts", len(subset) == sum(r['frozen_case_count'] for r in sample['records']))
    for row in sample['records']:
        b = subset[subset.requirement_id == row['requirement_id']]
        evidence.check(f"sample requirement {row['requirement_id']}", len(b) == row['frozen_case_count'] and set(b.verification_mode) == {row['verification_mode']} and set(b.source_family) == {row['source_family']})
    export(evidence.output, "model_sample", sample['records'])
    generation_identity = {}
    for model in MODELS:
        control = root / ("REFERENCE/LUNA/extraction_manifest.json" if model == 'luna' else f"OUTPUTS/FINAL/{model}/scaling_control/run_manifest.json")
        j = evidence.json(control, "selected model generation run", model)
        require(j is not None and j.get('status') == 'COMPLETE', f"Incomplete model generation control: {control}")
        evidence.check(f"{model} frozen sample hash", j.get('sample_manifest_sha256') == digest(root / 'MANIFESTS/sample_50.json'))
        if model != 'luna':
            evidence.check(f"{model} final generation scope", j.get('scope') == 'final' and j.get('completed_requirements') == 50)
        generation_identity[model] = j
        evidence.collect_references(j, control.parent, model)
    records, stage_run_ids = [], {}
    for stage, model in itertools.product(STAGES, MODELS):
        base = root / "OUTPUTS/RDF_EVALUATION" / stage / model
        identity = f"{model}_{stage}"
        manifest = evidence.json(base / 'run_manifest.json', 'selected model evaluation identity', identity)
        require(manifest is not None, f"Missing evaluation identity: {base}")
        expected_run = f"RDF-SCALING50-FINAL-{stage}-{model.upper()}"
        evidence.check(f"{identity} final evaluation run", manifest.get('run_id') == expected_run)
        rows = evidence.csv(base / 'raw_case_results.csv', 'model per-case outcomes', identity)
        require(all(r.get('run_id') == expected_run for r in rows), f"Mixed model evaluation runs: {base}")
        generation = genmap(evidence, root / 'RDF_INPUTS/FINAL' / stage / model / 'generated_rules.jsonl', identity)
        evidence.check(f"{identity} generation sample coverage", set(generation) == selected)
        stage_run_ids[(model, stage)] = {rid: r.get('run_id') for rid, r in generation.items()}
        records += normalise(evidence, rows, subset, model, stage, base, generation)
        evidence.collect_references(manifest, base, identity)
    for model in MODELS:
        evidence.check(f"{model} initial/final same generation runs", stage_run_ids[(model, STAGES[0])] == stage_run_ids[(model, STAGES[1])])
    return pd.DataFrame(records), selected


def score(frame, policy, exclusions, all_requirements):
    rows = frame[~frame.case_id.isin(exclusions)].copy()
    rows['policy'] = policy
    keys = ['policy', 'condition', 'stage', 'requirement_id', 'verification_mode', 'source_family']
    req = rows.groupby(keys, dropna=False).agg(cases=('case_id', 'size'), correct=('correct', 'sum'),
        verdicts=('observed', lambda s: int((s != 'NO_VERDICT').sum())),
        evaluation_rows=('evaluation_row_present', 'sum'), artifact_availability=('artifact_availability', 'first')).reset_index()
    req['requirement_accuracy'] = req.correct / req.cases
    req['all_cases_have_verdicts'] = req.verdicts == req.cases
    req['all_cases_correct'] = req.correct == req.cases
    summary, matrix = [], []
    for (condition, stage), group in rows.groupby(['condition', 'stage'], sort=False):
        scores = req[(req.condition == condition) & (req.stage == stage)]
        counts = Counter(group.scored_outcome)
        independently = sum(sum(g.correct) / len(g) for _, g in group.groupby('requirement_id')) / len(scores)
        require(math.isclose(float(scores.requirement_accuracy.mean()), independently, abs_tol=1e-12), 'Independent macro accuracy check failed')
        require(sum(counts.values()) == len(group), 'Outcome reconciliation failed')
        summary.append({'policy': policy, 'condition': condition, 'stage': stage, 'requirements': len(scores), 'cases': len(group),
            'requirement_level_accuracy': scores.requirement_accuracy.mean(), 'supplementary_pooled_case_accuracy': group.correct.mean(),
            'correct': counts['CORRECT'], 'false_accept': counts['FALSE_ACCEPT'], 'false_reject': counts['FALSE_REJECT'], 'no_verdict': counts['NO_VERDICT'],
            'requirements_with_all_verdicts': int(scores.all_cases_have_verdicts.sum()), 'exact_requirements': int(scores.all_cases_correct.sum()),
            'recorded_output_requirements': int((scores.artifact_availability == 'RECORDED_ARTIFACT').sum()),
            'recorded_no_output_requirements': int((scores.artifact_availability == 'RECORDED_NO_OUTPUT').sum()),
            'evaluation_complete': bool(group.evaluation_row_present.all()),
            'requirements_without_cases': '|'.join(sorted(all_requirements - set(scores.requirement_id)))})
        for expected, observed in itertools.product(['PASS', 'FAIL'], ['PASS', 'FAIL', 'NO_VERDICT']):
            n = int(((group.expected == expected) & (group.observed == observed)).sum())
            denominator = int((group.expected == expected).sum())
            matrix.append({'policy': policy, 'condition': condition, 'stage': stage, 'expected': expected, 'observed': observed,
                           'count': n, 'expected_row_total': denominator, 'row_fraction': n / denominator if denominator else None})
    return rows, req, pd.DataFrame(summary), pd.DataFrame(matrix)


def breakdowns(requirements):
    records = []
    for field in ['verification_mode', 'source_family']:
        for key, group in requirements.groupby(['policy', 'condition', 'stage', field]):
            records.append(dict(zip(['policy', 'condition', 'stage', 'category'], key)) | {'dimension': field,
                'requirements': len(group), 'cases': int(group.cases.sum()), 'requirement_level_accuracy': group.requirement_accuracy.mean(),
                'supplementary_pooled_case_accuracy': group.correct.sum() / group.cases.sum()})
    return pd.DataFrame(records)


def paired(requirements):
    records = []
    for policy, frame in requirements.groupby('policy'):
        workflows = frame[frame.condition.isin(WORKFLOWS)]
        for left, right in itertools.combinations(WORKFLOWS, 2):
            merged = workflows[workflows.condition == left].merge(workflows[workflows.condition == right], on='requirement_id', suffixes=('_left', '_right'), validate='one_to_one')
            for row in merged.to_dict('records'):
                records.append({'policy': policy, 'comparison_type': 'workflow', 'left': left, 'right': right,
                    'requirement_id': row['requirement_id'], 'left_accuracy': row['requirement_accuracy_left'], 'right_accuracy': row['requirement_accuracy_right'],
                    'difference_right_minus_left': row['requirement_accuracy_right'] - row['requirement_accuracy_left']})
        for model in MODELS:
            frame_m = frame[frame.condition == model]
            merged = frame_m[frame_m.stage == STAGES[0]].merge(frame_m[frame_m.stage == STAGES[1]], on='requirement_id', suffixes=('_left', '_right'), validate='one_to_one')
            for row in merged.to_dict('records'):
                records.append({'policy': policy, 'comparison_type': 'model', 'left': f'{model}/FIRST_GENERATION', 'right': f'{model}/SELF_REPAIR_FINAL',
                    'requirement_id': row['requirement_id'], 'left_accuracy': row['requirement_accuracy_left'], 'right_accuracy': row['requirement_accuracy_right'],
                    'difference_right_minus_left': row['requirement_accuracy_right'] - row['requirement_accuracy_left']})
    return pd.DataFrame(records)


def rubric_score(values, assessable):
    numeric, incomplete = [], []
    for value in values:
        text = str(value).strip().upper() if value is not None else ''
        if text in {'N/A', 'NA', 'NOT APPLICABLE'}:
            continue
        if text in {'0', '1', '2', '0.0', '1.0', '2.0'}:
            numeric.append(float(text))
        else:
            incomplete.append(text or 'BLANK')
    state = 'UNASSESSABLE' if not assessable else 'INCOMPLETE' if incomplete else 'NO_APPLICABLE_ITEMS' if not numeric else 'COMPLETE'
    return (sum(numeric) / (2 * len(numeric)) if state == 'COMPLETE' else None), len(numeric), state


def manual_data(evidence, sheets):
    scores, criteria, discrepancies = [], [], []
    for name in ['V2', 'NoSemantic', 'SingleShot', 'Scaling audit B01']:
        scaling = name == 'Scaling audit B01'
        rows = table(sheets[name], 12 if scaling else 6, 'Requirement ID')
        selected = rows if scaling else [r for r in rows if str(r.get('Manual review status', '')).upper() != 'NOT_REVIEWED']
        evidence.check(f'{name} selected manual sample', len(selected) == (80 if scaling else 70))
        for r in selected:
            condition = str(r.get('Model')) if scaling else {'V2': WORKFLOWS[-1], 'NoSemantic': 'NO_SEMANTIC', 'SingleShot': 'SINGLESHOT'}[name]
            stage = r.get('Evaluation mode', 'MANUAL_70')
            verdict = str(r.get('Primary semantic verdict') or 'UNCLEAR').strip().upper()
            require(verdict in VERDICTS, f'Unknown manual verdict {verdict}')
            status = r.get('Artifact status') if scaling else r.get('Manual review status')
            assessable = verdict != 'NOT_ASSESSABLE' and status not in {'NOT_ASSESSABLE', 'MISSING'}
            original = {'sheet': name, 'excel_row': r['_excel_row'], 'condition': condition, 'stage': stage,
                        'requirement_id': r['Requirement ID'], 'artifact_id': r.get('Artifact ID', ''),
                        'status': status, 'verdict': verdict, 'assessable': assessable,
                        'critical_defect': r.get('Critical defect', r.get('Critical defect?')),
                        'output_origin': r.get('Selected output origin', r.get('Artifact status')),
                        'original_fields': json.dumps(clean(r), ensure_ascii=False)}
            result = dict(original)
            for group, prefixes, cached_field in [('semantic', [f'Q{i}' for i in range(1, 8)], 'Semantic-fidelity normalized score' if scaling else 'Manual semantic-fidelity score'),
                                                   ('implementation', [f'CQ{i}' for i in range(1, 4)], 'Implementation-quality normalized score' if scaling else 'Implementation-quality score')]:
                values = []
                for prefix in prefixes:
                    key = next(k for k in r if isinstance(k, str) and k.startswith(prefix + ' ') and ('Score' in k or '(0–2)' in k))
                    value = r.get(key)
                    values.append(value)
                    criteria.append(original | {'group': group, 'criterion': prefix, 'raw_score': value,
                        'numeric_score': float(value) if str(value) in {'0', '1', '2', '0.0', '1.0', '2.0'} and assessable else None,
                        'evidence': r.get(prefix + ' Evidence', r.get('Code evidence / source / lines'))})
                recalculated, applicable, completion = rubric_score(values, assessable)
                result.update({f'{group}_score': recalculated, f'{group}_applicable_items': applicable, f'{group}_completion': completion,
                               f'{group}_cached_score': r.get(cached_field)})
                cached = r.get(cached_field)
                matches = (cached in (None, '') and recalculated is None) or (isinstance(cached, (float, int)) and recalculated is not None and math.isclose(cached, recalculated, abs_tol=1e-10))
                if not matches:
                    discrepancies.append(original | {'group': group, 'cached': cached, 'recalculated': recalculated, 'completion': completion,
                        'explanation': 'Recalculation excludes N/A and requires complete applicable criteria; unassessable outputs have no numeric score.'})
            scores.append(result)
    frame = pd.DataFrame(scores)
    summaries = []
    for key, group in frame.groupby(['condition', 'stage']):
        counts = Counter(group.verdict)
        summaries.append({'condition': key[0], 'stage': key[1], 'selected_outputs': len(group),
                          **{v: counts[v] for v in VERDICTS}, 'semantic_mean': group.semantic_score.mean(),
                          'semantic_mean_n': int(group.semantic_score.notna().sum()), 'implementation_mean': group.implementation_score.mean(),
                          'implementation_mean_n': int(group.implementation_score.notna().sum()),
                          'semantic_incomplete_n': int((group.semantic_completion == 'INCOMPLETE').sum()),
                          'implementation_incomplete_n': int((group.implementation_completion == 'INCOMPLETE').sum())})
        require(sum(counts.values()) == len(group), 'Manual categorical reconciliation failed')
    export(evidence.output, 'manual_scores', frame)
    export(evidence.output, 'manual_criteria', criteria)
    export(evidence.output, 'manual_summary', summaries)
    export(evidence.output, 'manual_cache_discrepancies', discrepancies)
    return frame, pd.DataFrame(criteria), pd.DataFrame(summaries), discrepancies


def formula_checks(evidence, sheets, requirements):
    discrepancies = []
    original = requirements[requirements.policy == 'ORIGINAL']
    index = {(r.condition, r.requirement_id): r.requirement_accuracy for r in original.itertuples() if r.condition in WORKFLOWS}
    for row in table(sheets['RequirementScores'], 6, 'Requirement ID'):
        condition = WORKFLOWS[-1] if row['Architecture'] == 'V2_FALLBACK25' else row['Architecture']
        value = index.get((condition, row['Requirement ID']))
        cached = row['Original requirement accuracy']
        if value is None or not isinstance(cached, (int, float)) or not math.isclose(cached, value, abs_tol=1e-10):
            discrepancies.append({'condition': condition, 'requirement_id': row['Requirement ID'], 'cached': cached, 'recalculated': value,
                                  'explanation': 'Original cached requirement accuracy compared against the frozen benchmark ledger.'})
    export(evidence.output, 'behavioural_cache_discrepancies', discrepancies)
    return discrepancies


def summary_cache_checks(evidence, sheets, cases, summary, manual_summary):
    checks = []
    mappings = {'NLTL_v2': WORKFLOWS[-1], 'Ablation case 1': 'NO_SEMANTIC', 'Ablation case 2': 'SINGLESHOT',
                'GPT-5.6 Luna': 'luna', 'GPT-5.6 SOL': 'sol', 'Gemini 3.5 Flash': 'gemini', 'gpt-oss-120b': 'gpt_oss'}
    def compare(sheet, row_number, col, metric, cached, recalculated, explanation):
        match = (cached in (None, '') and pd.isna(recalculated)) or (isinstance(cached, (int, float)) and not pd.isna(recalculated) and math.isclose(cached, float(recalculated), abs_tol=1e-10))
        checks.append({'sheet': sheet, 'cell': f'{openpyxl.utils.get_column_letter(col)}{row_number}',
                       'metric': metric, 'cached': cached, 'recalculated': recalculated, 'matches': match, 'explanation': explanation})
    for sheet in ['Paper stats 1', 'Paper stats 2', 'Paper stats 3']:
        stage = 'RUN_01' if sheet == 'Paper stats 1' else STAGES[1] if sheet == 'Paper stats 2' else STAGES[0]
        count = 3 if sheet == 'Paper stats 1' else 4
        for number in range(7, 7 + count):
            row = sheets[sheet][number - 1]
            condition = mappings[row[0]]
            data = cases[(cases.policy == 'ORIGINAL') & (cases.condition == condition) & (cases.stage == stage)]
            s = summary[(summary.policy == 'ORIGINAL') & (summary.condition == condition) & (summary.stage == stage)].iloc[0]
            values = [int(((data.expected == 'PASS') & (data.observed == 'PASS')).sum()),
                      int(((data.expected == 'FAIL') & (data.observed == 'FAIL')).sum()), s.false_accept, s.false_reject,
                      s.correct, s.cases, s.supplementary_pooled_case_accuracy]
            for col, value in enumerate(values, 2):
                compare(sheet, number, col, sheets[sheet][5][col - 1], row[col - 1], value, 'Original benchmark; behavioural accuracy in this sheet is pooled case accuracy.')
        for number in range(23, 23 + count):
            row = sheets[sheet][number - 1]
            label = mappings[row[0]] if sheet == 'Paper stats 1' else row[0]
            manual_stage = 'MANUAL_70' if sheet == 'Paper stats 1' else stage
            group = manual_summary[(manual_summary.condition == label) & (manual_summary.stage == manual_stage)]
            require(len(group) == 1, f'Manual summary identity mismatch: {sheet}/{label}')
            s = group.iloc[0]
            values = [s.selected_outputs, s.CORRECT, s.CORRECT / s.selected_outputs,
                      s.MINOR_DEFECT, s.MINOR_DEFECT / s.selected_outputs, s.MAJOR_DEFECT, s.MAJOR_DEFECT / s.selected_outputs,
                      s.NOT_ASSESSABLE, s.NOT_ASSESSABLE / s.selected_outputs, s.UNCLEAR, s.UNCLEAR / s.selected_outputs,
                      s.semantic_mean, s.implementation_mean]
            for col, value in enumerate(values, 2):
                compare(sheet, number, col, sheets[sheet][21][col - 1], row[col - 1], value,
                        'All selected slots are retained in categorical denominators; numeric means use only complete assessable scores.')
    frame = pd.DataFrame(checks)
    export(evidence.output, 'summary_cache_checks', frame)
    return frame[~frame.matches].to_dict('records')


def rubric_heatmap(data):
    """Keep each assessment distinct; blinded slot IDs repeat per requirement."""
    temp = data.copy()
    require(not temp.duplicated(['excel_row', 'criterion']).any(),
            'Duplicate rubric criterion within an assessment worksheet row')
    slots = temp.drop_duplicates('excel_row').set_index('excel_row')
    labels = {}
    for row_number, slot in slots.iterrows():
        artifact_id = slot.get('artifact_id')
        label = str(slot['requirement_id'])
        if pd.notna(artifact_id) and str(artifact_id).strip():
            label += ' / ' + str(artifact_id).strip()
        labels[row_number] = label
    heat = temp.pivot(index='excel_row', columns='criterion', values='numeric_score')
    heat = heat.reindex(columns=[f'Q{i}' for i in range(1, 8)] + [f'CQ{i}' for i in range(1, 4)])
    heat.index = pd.Index([labels[row_number] for row_number in heat.index], name='assessment')
    return heat


def save_figure(fig, directory, name, inventory):
    """Use constrained layout and tight export; save all publication formats."""
    fig.canvas.draw()
    # Renderer checks detect text extending beyond the figure canvas. Tight export
    # accommodates external legends; a human should still review publication choices.
    renderer = fig.canvas.get_renderer()
    bounds = fig.bbox
    outside = []
    for text in fig.findobj(matplotlib.text.Text):
        if text.get_visible() and text.get_text():
            box = text.get_window_extent(renderer)
            if box.x0 < bounds.x0 - 2 or box.y0 < bounds.y0 - 2 or box.x1 > bounds.x1 + 2 or box.y1 > bounds.y1 + 2:
                outside.append(text.get_text())
    for suffix in ['pdf', 'svg', 'png']:
        fig.savefig(directory / f'{name}.{suffix}', dpi=300, bbox_inches='tight', pad_inches=0.12)
    inventory.append({'figure': name, 'external_text_before_tight_export': '|'.join(outside), 'formats': 'PDF|SVG|PNG', 'dpi': 300})
    plt.close(fig)


def figures(output, summary, matrices, breakdown, requirements, manual, criteria, manual_summary, failure_counts):
    directory = output / 'figures'
    directory.mkdir()
    inventory = []
    plt.rcParams.update({'font.size': 10, 'axes.labelsize': 10, 'legend.fontsize': 9, 'svg.fonttype': 'none', 'pdf.fonttype': 42})
    for policy, subset in summary.groupby('policy', sort=False):
        stem = policy.lower()
        workflow = subset[subset.condition.isin(WORKFLOWS)].set_index('condition').reindex(WORKFLOWS)
        fig, axes = plt.subplots(1, 2, figsize=(10, 4), layout='constrained', gridspec_kw={'width_ratios': [3, 1]})
        for ax, names in zip(axes, [WORKFLOWS[:3], WORKFLOWS[3:]]):
            values = workflow.loc[names, 'requirement_level_accuracy'] * 100
            bars = ax.bar(range(len(names)), values, color=[COLORS[n] for n in names])
            ax.bar_label(bars, fmt='%.1f%%', padding=3)
            ax.set_xticks(range(len(names)), ['Full repair V2' if n == WORKFLOWS[0] else 'No semantic' if n == WORKFLOWS[1] else 'Single shot' if n == WORKFLOWS[2] else 'V2 + rejected 25' for n in names], rotation=15)
            ax.set_ylim(0, 108)
            ax.set_yticks(range(0, 101, 20))
            ax.set_title('Official workflows' if len(names) == 3 else 'Retrospective')
        axes[0].set_ylabel('Requirement accuracy (%)')
        save_figure(fig, directory, f'{stem}_workflow_accuracy', inventory)
        fig, axes = plt.subplots(2, 2, figsize=(10, 7), layout='constrained')
        for ax, condition in zip(axes.flat, WORKFLOWS):
            data = matrices[(matrices.policy == policy) & (matrices.condition == condition)]
            counts = data.pivot(index='expected', columns='observed', values='count').reindex(index=['PASS', 'FAIL'], columns=['PASS', 'FAIL', 'NO_VERDICT']).to_numpy()
            frac = data.pivot(index='expected', columns='observed', values='row_fraction').reindex(index=['PASS', 'FAIL'], columns=['PASS', 'FAIL', 'NO_VERDICT']).to_numpy()
            ax.imshow(frac, vmin=0, vmax=1, cmap='Blues', aspect='auto')
            for i, j in itertools.product(range(2), range(3)):
                ax.text(j, i, f'{counts[i,j]:,}\n{frac[i,j]:.1%}', ha='center', va='center', color='white' if frac[i,j] > .55 else 'black')
            ax.set_xticks(range(3), ['PASS', 'FAIL', 'No verdict'])
            ax.set_yticks(range(2), ['PASS', 'FAIL'])
            ax.set_ylabel('Expected')
            ax.set_xlabel('Observed')
            ax.set_title(condition.replace('_', ' '), fontsize=9)
        save_figure(fig, directory, f'{stem}_workflow_outcomes', inventory)
        model = subset[subset.condition.isin(MODELS)]
        fig, ax = plt.subplots(figsize=(7, 4), layout='constrained')
        for m in MODELS:
            vals = model[model.condition == m].set_index('stage').reindex(STAGES)
            ax.plot([0, 1], vals.requirement_level_accuracy * 100, marker='o', color=COLORS[m], label=f'{m} (n={int(vals.requirements.iloc[0])})')
        ax.set_xticks([0, 1], ['First generation', 'Self-repair final'])
        ax.set_ylim(0, 100)
        ax.set_xlim(-.12, 1.12)
        ax.set_ylabel('Requirement accuracy (%)')
        ax.legend(loc='center left', bbox_to_anchor=(1.02, .5))
        save_figure(fig, directory, f'{stem}_model_progression', inventory)
        fig, ax = plt.subplots(figsize=(7, 3.5), layout='constrained')
        for i, m in enumerate(MODELS):
            vals = model[model.condition == m].set_index('stage').reindex(STAGES).requirement_level_accuracy.to_numpy() * 100
            ax.plot(vals, [i, i], color=COLORS[m], linewidth=2)
            ax.scatter(vals[0], i, color=COLORS[m], marker='o', label='First generation' if i == 0 else None)
            ax.scatter(vals[1], i, color=COLORS[m], marker='D', label='Self-repair final' if i == 0 else None)
        ax.set_yticks(range(4), MODELS)
        ax.set_xlim(0, 100)
        ax.set_xlabel('Requirement accuracy (%)')
        ax.legend(loc='center left', bbox_to_anchor=(1.02, .5))
        save_figure(fig, directory, f'{stem}_model_dumbbell', inventory)
        data = breakdown[(breakdown.policy == policy) & (breakdown.dimension == 'verification_mode') & breakdown.condition.isin(MODELS)]
        modes = sorted(data.category.unique())
        fig, axes = plt.subplots(1, len(modes), figsize=(12, 4), layout='constrained', squeeze=False)
        for ax, mode in zip(axes.flat, modes):
            for m in MODELS:
                vals = data[(data.condition == m) & (data.category == mode)].set_index('stage').reindex(STAGES)
                ax.plot([0, 1], vals.requirement_level_accuracy * 100, 'o-', color=COLORS[m], label=m)
            ax.set_title(f'{mode.replace("_", " ")}\nn={int(vals.requirements.iloc[0])}', fontsize=9)
            ax.set_xticks([0, 1], ['First', 'Final'])
            ax.set_ylim(0, 100)
        axes.flat[0].set_ylabel('Requirement accuracy (%)')
        axes.flat[-1].legend(loc='center left', bbox_to_anchor=(1.02, .5))
        save_figure(fig, directory, f'{stem}_model_by_mode', inventory)
        data = requirements[(requirements.policy == policy) & requirements.condition.isin(MODELS)]
        heat = data.pivot(index='requirement_id', columns=['condition', 'stage'], values='requirement_accuracy')
        heat = heat.reindex(columns=pd.MultiIndex.from_product([MODELS, STAGES]))
        fig, ax = plt.subplots(figsize=(10, max(9, len(heat) * .19)), layout='constrained')
        im = ax.imshow(heat.to_numpy(), vmin=0, vmax=1, aspect='auto', cmap='viridis')
        ax.set_yticks(range(len(heat)), heat.index, fontsize=7)
        ax.set_xticks(range(8), [f'{m}\n{"First" if s == STAGES[0] else "Final"}' for m, s in heat.columns], fontsize=8)
        fig.colorbar(im, ax=ax, label='Requirement accuracy')
        save_figure(fig, directory, f'{stem}_model_requirements', inventory)
    groups = manual_summary.copy()
    labels = [f'{r.condition}\n{r.stage}' for r in groups.itertuples()]
    fig, ax = plt.subplots(figsize=(11, 7), layout='constrained')
    left = np.zeros(len(groups))
    for verdict, color in zip(VERDICTS, VERDICT_COLORS):
        ax.barh(range(len(groups)), groups[verdict], left=left, label=verdict.replace('_', ' '), color=color)
        left += groups[verdict].to_numpy()
    ax.set_yticks(range(len(groups)), labels, fontsize=8)
    ax.set_xlabel('Reviewed output slots')
    ax.legend(loc='center left', bbox_to_anchor=(1.02, .5))
    save_figure(fig, directory, 'manual_verdict_distributions', inventory)
    for field in ['semantic', 'implementation']:
        fig, ax = plt.subplots(figsize=(11, 7), layout='constrained')
        positions = []
        for i, r in enumerate(groups.itertuples()):
            values = manual[(manual.condition == r.condition) & (manual.stage == r.stage)][field + '_score'].dropna()
            if len(values):
                ax.boxplot([values.to_numpy() * 100], positions=[i], widths=.5, orientation='horizontal', manage_ticks=False)
                ax.scatter(values * 100, np.full(len(values), i), s=12, alpha=.5, color='#0072B2')
            positions.append(f'{r.condition}\n{r.stage} (n={len(values)})')
        ax.set_yticks(range(len(groups)), positions, fontsize=8)
        ax.set_xlim(-2, 102)
        ax.set_xlabel(field.capitalize() + ' score (%)')
        save_figure(fig, directory, f'manual_{field}_scores', inventory)
    for name, data in criteria.groupby('sheet'):
        heat = rubric_heatmap(data)
        fig, ax = plt.subplots(figsize=(9, max(8, len(heat) * .16)), layout='constrained')
        cmap = plt.get_cmap('viridis').copy()
        cmap.set_bad('#dddddd')
        im = ax.imshow(heat.to_numpy(dtype=float), cmap=cmap, vmin=0, vmax=2, aspect='auto')
        ax.set_xticks(range(10), heat.columns, fontsize=9)
        ax.set_yticks(range(len(heat)), heat.index, fontsize=6)
        fig.colorbar(im, ax=ax, ticks=[0, 1, 2], label='Criterion score')
        safe_name = re.sub(r'\W', '_', name)
        save_figure(fig, directory, f'manual_rubric_{safe_name}', inventory)
    # Exact recorded failure stages are comparable; no inferred ontology/syntax taxonomy.
    for policy, group in failure_counts.groupby('policy'):
        pivot = group.pivot_table(index='failure_reason_evidence_label', columns='condition', values='cases', aggfunc='sum', fill_value=0)
        if not pivot.empty:
            fig, ax = plt.subplots(figsize=(12, max(4, len(pivot) * .45)), layout='constrained')
            pivot.plot.barh(ax=ax, color=[COLORS.get(c, '#777777') for c in pivot.columns])
            ax.set_xlabel('No-verdict cases')
            ax.set_ylabel('Recorded diagnostic label')
            ax.legend(loc='center left', bbox_to_anchor=(1.02, .5), fontsize=8)
            save_figure(fig, directory, policy.lower() + '_failure_labels', inventory)
    export(output, 'figure_verification', inventory)
    return inventory


def make_reports(output, summary, breakdown, manual_summary, exclusions, contribution, discrepancies, incomplete, inventory):
    status = 'INCOMPLETE — diagnostic report only' if incomplete else 'Complete saved evaluation coverage'
    text = ['# Ocean Engineering results', '', status, '',
        'Equal-weight requirement accuracy averages correct/included-case fractions across requirements with at least one included case. Missing verdicts earn zero. Pooled case accuracy is supplementary. PASS is the positive class.', '',
        'Original frozen labels are used throughout. RDF_EDIT sensitivity excludes proposed edits, not confirmed invalid cases. Explicit exclusions are a separate scenario and are not automatically unioned with proposed edits.', '',
        'Official FULL_REPAIR_V2 RUN_01 is separate from the retrospective substitution of 25 rejected fourth-attempt candidates. Historical FULL is excluded. Single selected runs do not support significance claims or general scaling laws.', '',
        'Workflow manual review covers 70 selected requirements per condition; model review covers ten per model/stage. All manual subset figures are exploratory. The V2 manual review describes the retrospective selection. Semantic and implementation scores are distinct; a major defect remains a major defect regardless of mean score.', '',
        'N/A is excluded from rubric denominators; blank/UNCLEAR prevents a final numeric score. Unassessable slots remain in verdict counts and are outside numeric means.', '',
        'NOT_PARSED alone is not evidence of invalid Turtle. Failure comparisons use recorded stage/exception labels; behavioural mismatch causes remain undetermined unless the diagnostics establish them.', '',
        '## Behavioural results', '', '```csv', summary.to_csv(index=False), '```', '',
        '## Exclusion policy and sample sizes', '', exclusions.to_csv(index=False), '',
        '## Manual review', '', manual_summary.to_csv(index=False), '',
        '## Retrospective contribution', '', contribution.to_csv(index=False), '',
        '## Cached workbook checks', '', f'{len(discrepancies)} cached score discrepancies. See manual_cache_discrepancies and behavioural_cache_discrepancies. Original formulas and caches are exported; no formula recalculation or workbook modification is performed.', '',
        'Paper-summary outcome counts, accuracy, and manual categorical counts/means are checked in summary_cache_checks.csv. The manual comparison uses all selected slots, so reviewed-only workbook summaries may omit unassessable slots. These denominator differences are reported explicitly. All original formulas/caches are preserved.', '',
        '## Verification and figures', '', 'Outcome totals and independent macro means were reconciled during this run. Frozen IDs, stage generation-run alignment and exclusion policies were checked. Full details: verification.csv.', '',
        'Figure exports use constrained layout and tight bounds. figure_verification.csv records text outside the initial canvas. Review the HTML gallery visually before manuscript submission; automated bounds checks cannot prove absence of all overlapping text.', '',
        'See sources_used.csv for hashed inputs and referenced_files.csv for evidence availability. See MISSING_DATA.md for missing records.', '']
    for item in inventory:
        text += [f"![{item['figure']}](figures/{item['figure']}.png)", '']
    (output / 'RESULTS_REPORT.md').write_text('\n'.join(text), encoding='utf-8')
    gallery = ''.join(f'<figure><a href="figures/{i["figure"]}.svg"><img src="figures/{i["figure"]}.png" alt="{i["figure"]}" loading="lazy"></a><figcaption>{i["figure"]} · <a href="figures/{i["figure"]}.pdf">PDF</a></figcaption></figure>' for i in inventory)
    page = f'''<!doctype html><html lang="en"><meta charset="utf-8"><title>Ocean Engineering results</title>
<style>body{{font:15px system-ui;margin:2rem;line-height:1.5}}table{{border-collapse:collapse;font-size:12px;display:block;overflow:auto}}th,td{{padding:.4rem;border:1px solid #ddd}}img{{width:100%;max-width:1100px}}figure{{margin:2rem 0}}pre{{white-space:pre-wrap}}</style>
<h1>Ocean Engineering results</h1><p>{html.escape(status)}</p><p>Manual subset figures are exploratory. Requirement accuracy uses equal requirement weights; pooled case accuracy is supplementary. Retrospective candidates remain separate from official outputs.</p>
<nav><a href="RESULTS_REPORT.md">Methods and limitations</a> · <a href="MISSING_DATA.md">Missing data</a> · <a href="sources_used.csv">Sources</a> · <a href="verification.csv">Verification</a></nav>
<h2>Behavioural results</h2>{summary.to_html(index=False, escape=True)}<h2>Exclusions</h2>{exclusions.to_html(index=False, escape=True)}
<h2>Manual review</h2>{manual_summary.to_html(index=False, escape=True)}<h2>Retrospective contribution</h2>{contribution.to_html(index=False, escape=True)}
<h2>Figures</h2>{gallery}</html>'''
    (output / 'report.html').write_text(page, encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--repo-root', type=Path, default=HERE.parent)
    parser.add_argument('--workbook', type=Path, default=HERE / 'inputs/NLTL_Manual_Audit_with_Model_Scaling_FINAL (1) (4).xlsx')
    parser.add_argument('--output-dir', type=Path, default=HERE / 'outputs/report')
    parser.add_argument('--exclude-ids', type=Path, help='UTF-8 text, one case ID per line, or CSV with case_id header')
    parser.add_argument('--input-config', type=Path, help='JSON overrides: official, historical, rejected evaluation directories and scaling_root')
    parser.add_argument('--allow-incomplete', action='store_true', help='Produce explicitly incomplete diagnostic results; never claim complete experiment coverage')
    parser.add_argument('--no-figures', action='store_true', help='Export data/reports only')
    args = parser.parse_args()
    root, output, workbook = args.repo_root.resolve(), args.output_dir.resolve(), args.workbook.resolve()
    require(not output.exists(), f'Output already exists: {output}. Choose a new --output-dir to preserve prior reports.')
    require(not workbook.is_relative_to(output), 'Workbook must not be inside report output')
    output.mkdir(parents=True)
    evidence = Evidence(root, output)
    try:
        overrides = evidence.json(args.input_config.resolve(), 'explicit input overrides', 'report configuration') if args.input_config else {}
        require(isinstance(overrides, dict), 'Input configuration must be an object')
        # All override paths must be explicit absolute paths: no working-directory ambiguity.
        require(set(overrides) <= {'official', 'historical', 'rejected', 'scaling_root'}, 'Unknown input-config field')
        require(all(Path(v).is_absolute() for v in overrides.values()), 'Override directories must be absolute paths')
        write_json(output / 'execution_config.json', vars(args) | {'repo_root': str(root), 'workbook': str(workbook), 'output_dir': str(output), 'overrides': overrides})
        before = digest(workbook)
        sheets = workbook_data(evidence, workbook)
        benchmark, proposed = load_benchmark(evidence, sheets)
        export(output, 'benchmark', benchmark)
        explicit = set()
        if args.exclude_ids:
            path = args.exclude_ids.resolve()
            require(evidence.use(path, 'explicit confirmed/user exclusion IDs', 'EXPLICIT_EXCLUSIONS'), f'Exclusion file missing: {path}')
            if path.suffix.lower() == '.csv':
                with path.open(encoding='utf-8-sig', newline='') as stream:
                    reader = csv.DictReader(stream)
                    require('case_id' in (reader.fieldnames or []), 'Exclusion CSV needs case_id header')
                    explicit = {r['case_id'].strip() for r in reader if r['case_id'].strip()}
            else:
                explicit = {line.strip() for line in path.read_text(encoding='utf-8-sig').splitlines() if line.strip() and not line.lstrip().startswith('#')}
            require(explicit <= set(benchmark.case_id), f'Unknown exclusion IDs: {explicit - set(benchmark.case_id)}')
        export(output, 'explicit_exclusions', [{'case_id': cid} for cid in sorted(explicit)])
        workflow = load_workflows(evidence, benchmark, selected_runs(evidence, overrides))
        models, model_ids = load_models(evidence, benchmark, overrides)
        incomplete = not workflow.evaluation_row_present.all() or not models.evaluation_row_present.all()
        export(output, 'workflow_case_outcomes', workflow)
        export(output, 'model_case_outcomes', models)
        if incomplete and not args.allow_incomplete:
            raise ValueError('Incomplete evaluation files: see MISSING_DATA.md. Restore the precise saved outcome rows or explicitly request --allow-incomplete diagnostic reporting.')
        policies = {'ORIGINAL': set(), 'PROPOSED_RDF_EDIT_SENSITIVITY': proposed}
        if args.exclude_ids:
            policies['EXPLICIT_EXCLUSIONS'] = explicit
        summary_frames, req_frames, case_frames, matrix_frames, exclusion_rows = [], [], [], [], []
        for policy, excluded in policies.items():
            for kind, frame, ids in [('workflow', workflow, set(benchmark.requirement_id)), ('model', models, model_ids)]:
                rows, req, summary, matrix = score(frame, policy, excluded, ids)
                case_frames.append(rows)
                req_frames.append(req)
                summary_frames.append(summary)
                matrix_frames.append(matrix)
                included = benchmark[benchmark.requirement_id.isin(ids) & ~benchmark.case_id.isin(excluded)]
                overlap = set(benchmark[benchmark.requirement_id.isin(ids)].case_id) & excluded
                exclusion_rows.append({'policy': policy, 'experiment': kind, 'excluded_cases': len(overlap), 'included_cases': len(included),
                    'included_requirements': included.requirement_id.nunique(), 'requirements_without_cases': '|'.join(sorted(ids - set(included.requirement_id))),
                    'excluded_case_ids': '|'.join(sorted(overlap))})
                for _, group in rows.groupby(['condition', 'stage']):
                    evidence.check(f'{policy} {group.condition.iloc[0]} {group.stage.iloc[0]} exclusions identical', set(group.case_id) == set(included.case_id))
        summary, requirements = pd.concat(summary_frames, ignore_index=True), pd.concat(req_frames, ignore_index=True)
        cases, matrices = pd.concat(case_frames, ignore_index=True), pd.concat(matrix_frames, ignore_index=True)
        breakdown = breakdowns(requirements)
        exclusions = pd.DataFrame(exclusion_rows)
        export(output, 'policy_case_outcomes', cases)
        export(output, 'requirement_scores', requirements)
        export(output, 'summary_results', summary)
        export(output, 'outcome_matrices', matrices)
        export(output, 'breakdowns', breakdown)
        export(output, 'paired_differences', paired(requirements))
        export(output, 'exclusion_policy_summary', exclusions)
        failures = cases[cases.scored_outcome != 'CORRECT']
        export(output, 'failure_diagnostics_cases', failures)
        failure_counts = cases[cases.observed == 'NO_VERDICT'].groupby(['policy', 'condition', 'stage', 'failure_reason_evidence_label']).agg(cases=('case_id', 'size'), requirements=('requirement_id', 'nunique')).reset_index()
        export(output, 'failure_reason_counts', failure_counts)
        export(output, 'failure_diagnostics_requirements', failures.groupby(['policy', 'condition', 'stage', 'requirement_id', 'failure_reason_evidence_label']).agg(affected_cases=('case_id', 'size'), case_ids=('case_id', lambda s: '|'.join(sorted(s)))).reset_index())
        contribution = []
        for policy, rows in cases.groupby('policy'):
            official = rows[rows.condition == WORKFLOWS[0]].set_index('case_id')
            retro = rows[(rows.condition == WORKFLOWS[-1]) & rows.retrospective_candidate].set_index('case_id')
            contribution.append({'policy': policy, 'candidate_requirements': retro.requirement_id.nunique(), 'candidate_cases': len(retro),
                'candidate_correct': int(retro.correct.sum()), 'net_correct_gain': int(retro.correct.sum() - official.loc[retro.index].correct.sum()),
                'official_accuracy': summary[(summary.policy == policy) & (summary.condition == WORKFLOWS[0])].requirement_level_accuracy.iloc[0],
                'retrospective_accuracy': summary[(summary.policy == policy) & (summary.condition == WORKFLOWS[-1])].requirement_level_accuracy.iloc[0]})
        contribution = pd.DataFrame(contribution)
        export(output, 'retrospective_contribution', contribution)
        manual, criteria, manual_summary, manual_discrepancies = manual_data(evidence, sheets)
        cache_discrepancies = formula_checks(evidence, sheets, requirements)
        summary_discrepancies = summary_cache_checks(evidence, sheets, cases, summary, manual_summary)
        # Hash configuration and selection metadata without exporting raw config contents.
        for path in [root / SCALING / 'MANIFESTS/configuration_lock.json', root / SCALING / 'MANIFESTS/manifest_lock.json',
                     root / SCALING / 'MANUAL_AUDIT_SCALING_10/selection_manifest.json']:
            evidence.use(path, 'configuration/selection provenance', 'SCALING50', required=False)
        for exp, filename in [('FULL_REPAIR_V2', 'pipeline.full-repair-v2-run01.json'), ('LUNA_NO_SEMANTIC_VALIDATOR', 'pipeline.luna-no-semantic-validator-run01.json'), ('LUNA_CONTEXTUAL_SINGLESHOT', 'pipeline.luna-contextual-singleshot-run01.json')]:
            evidence.use(root / PIPE / 'experiments' / exp / 'CONFIGS' / filename, 'workflow configuration', exp, required=False)
        evidence.check('source workbook unchanged', digest(workbook) == before)
        evidence.check('complete evaluation rows', not incomplete or args.allow_incomplete, 'Incomplete diagnostic mode' if incomplete else 'Complete')
        inventory = [] if args.no_figures else figures(output, summary, matrices, breakdown, requirements, manual, criteria, manual_summary, failure_counts)
        make_reports(output, summary, breakdown, manual_summary, exclusions, contribution, manual_discrepancies + cache_discrepancies + summary_discrepancies, incomplete, inventory)
        evidence.save()
        # Only explicitly generated reports/data/figures; never recursive repository packaging.
        zip_path = output.parent / (output.name + '.zip')
        require(not zip_path.exists(), f'ZIP already exists: {zip_path}')
        with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as archive:
            for path in sorted(output.rglob('*')):
                if path.is_file():
                    archive.write(path, Path(output.name) / path.relative_to(output))
        print(summary.to_string(index=False))
        print(f'Output directory: {output}\nShareable ZIP: {zip_path}')
    except Exception as error:
        evidence.save()
        (output / 'RUN_FAILED.md').write_text(f'# Report did not complete\n\n{type(error).__name__}: {error}\n\nPartial exports are diagnostic only. No completed results ZIP was produced.\n', encoding='utf-8')
        print(f'Report failed: {error}\nDetails: {output / "MISSING_DATA.md"}', file=sys.stderr)
        raise


if __name__ == '__main__':
    main()
