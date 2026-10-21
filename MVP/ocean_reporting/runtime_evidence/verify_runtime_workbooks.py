from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


ROOT = Path("/Users/sadisfaction570/Desktop/Journal 1/NLTL_v2/MVP/ocean_reporting")
RUNTIME = ROOT / "NLTL_Runtime_and_Iterations.xlsx"
PAPER_SOURCE = Path("/Users/sadisfaction570/Downloads/NLTL_v2_Paper_Tables.xlsx")
PAPER_COPY = ROOT / "inputs/NLTL_v2_Paper_Tables.xlsx"


def assert_zip_ok(path: Path) -> None:
    with ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            raise AssertionError(f"Corrupt member in {path}: {bad}")


def compare_preserved_sheets() -> list[str]:
    source = load_workbook(PAPER_SOURCE, data_only=False, read_only=False)
    copy = load_workbook(PAPER_COPY, data_only=False, read_only=False)
    if copy.sheetnames[:-1] != source.sheetnames:
        raise AssertionError((source.sheetnames, copy.sheetnames))
    if copy.sheetnames[-1] != "Runtime summary":
        raise AssertionError(copy.sheetnames)
    for sheet_name in source.sheetnames:
        left = source[sheet_name]
        right = copy[sheet_name]
        if left.max_row != right.max_row or left.max_column != right.max_column:
            raise AssertionError(
                f"Dimension changed for {sheet_name}: "
                f"{left.max_row}x{left.max_column} -> {right.max_row}x{right.max_column}"
            )
        for row in left.iter_rows():
            for cell in row:
                if cell.value != right[cell.coordinate].value:
                    raise AssertionError(
                        f"Value/formula changed in {sheet_name}!{cell.coordinate}: "
                        f"{cell.value!r} -> {right[cell.coordinate].value!r}"
                    )
    return source.sheetnames


def verify_runtime() -> dict[str, tuple[int, int]]:
    workbook = load_workbook(RUNTIME, data_only=False, read_only=False)
    expected = [
        "Summary",
        "README",
        "Experiment_Index",
        "Requirement_Runs",
        "Iterations",
        "API_Calls",
        "RDF_Evaluation",
        "Audit",
    ]
    if workbook.sheetnames != expected:
        raise AssertionError(workbook.sheetnames)
    dimensions = {
        name: (workbook[name].max_row, workbook[name].max_column)
        for name in expected
    }
    required_data_rows = {
        "Experiment_Index": 25,
        "Requirement_Runs": 6046,
        "Iterations": 8849,
        "API_Calls": 17482,
        "RDF_Evaluation": 11792,
        "Audit": 11,
    }
    for name, count in required_data_rows.items():
        actual = dimensions[name][0] - 4
        if actual != count:
            raise AssertionError(f"{name}: expected {count} data rows, found {actual}")
    errors = []
    error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(token in cell.value for token in error_tokens):
                    errors.append(f"{sheet.title}!{cell.coordinate}={cell.value}")
    if errors:
        raise AssertionError(errors[:10])
    return dimensions


if __name__ == "__main__":
    assert_zip_ok(RUNTIME)
    assert_zip_ok(PAPER_COPY)
    preserved = compare_preserved_sheets()
    dimensions = verify_runtime()
    print("ZIP_INTEGRITY=OK")
    print(f"PRESERVED_PAPER_SHEETS={len(preserved)}")
    print("PAPER_EXISTING_VALUES_AND_FORMULAS=UNCHANGED")
    print("RUNTIME_SHEETS=" + ",".join(dimensions))
    for name, shape in dimensions.items():
        print(f"{name}={shape[0]}x{shape[1]}")
    print("FORMULA_ERROR_SCAN=OK")
